from __future__ import annotations

import re
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "PROJECT_REPORT.md"
OUT_DIR = ROOT / "docs" / "out"
DOCX_OUT = OUT_DIR / "心潮手帐项目报告.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
NAVY = RGBColor(32, 55, 72)
MUTED = RGBColor(98, 108, 116)
LIGHT_FILL = "F4F6F9"
TABLE_FILL = "F2F4F7"
WHITE = RGBColor(255, 255, 255)


def set_run_font(run, name="Microsoft YaHei", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_width(table, widths):
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("第 ")
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = "PAGE"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.append(fld_begin)
    run._r.append(instr)
    run._r.append(fld_end)
    paragraph.add_run(" 页")


def configure_document(doc: Document):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.2

    code_style = styles.add_style("CodeBlock", 1)
    code_style.font.name = "Consolas"
    code_style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    code_style.font.size = Pt(8.5)
    code_style.paragraph_format.space_before = Pt(0)
    code_style.paragraph_format.space_after = Pt(0)
    code_style.paragraph_format.line_spacing = 1.05

    header = section.header.paragraphs[0]
    header.text = "心潮手帐 | 智能情绪陪伴系统项目报告"
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    set_run_font(header.runs[0], size=9, color=MUTED)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)
    for run in footer.runs:
        set_run_font(run, size=9, color=MUTED)


def add_cover(doc: Document):
    for _ in range(5):
        doc.add_paragraph()

    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = kicker.add_run("项目报告")
    set_run_font(r, size=12, color=BLUE, bold=True)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)
    r = title.add_run("《心潮手帐》")
    set_run_font(r, size=30, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(26)
    r = subtitle.add_run("智能情绪陪伴系统")
    set_run_font(r, size=17, color=DARK_BLUE)

    summary = doc.add_paragraph()
    summary.alignment = WD_ALIGN_PARAGRAPH.CENTER
    summary.paragraph_format.left_indent = Inches(0.45)
    summary.paragraph_format.right_indent = Inches(0.45)
    summary.paragraph_format.space_after = Pt(30)
    r = summary.add_run("融合情绪识别、阅读推荐、AI 文学陪伴与可视化手账创作的全栈智能应用")
    set_run_font(r, size=11.5, color=MUTED, italic=True)

    table = doc.add_table(rows=4, cols=2)
    table.style = "Table Grid"
    set_table_width(table, [1.6, 4.4])
    rows = [
        ("系统名称", "心潮手帐 Heartide Journal"),
        ("项目类型", "智能情绪记录、阅读疗愈与手账创作系统"),
        ("技术方向", "机器学习 + 大模型生成 + 推荐系统 + 前后端工程"),
        ("生成日期", "2026 年 6 月 16 日"),
    ]
    for i, (label, value) in enumerate(rows):
        shade_cell(table.cell(i, 0), TABLE_FILL)
        for j, text in enumerate([label, value]):
            p = table.cell(i, j).paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(text)
            set_run_font(run, size=10.5, color=NAVY if j == 0 else RGBColor(40, 40, 40), bold=(j == 0))

    doc.add_page_break()


def add_static_toc(doc: Document):
    doc.add_heading("目录", level=1)
    items = [
        "一、项目概述",
        "二、系统技术栈",
        "三、系统架构",
        "四、主要功能模块",
        "五、情绪模型与训练数据",
        "六、Prompt 设计",
        "七、生图模型 API 接入状态",
        "八、安全与生产化设计",
        "九、测试与验证",
        "十、项目亮点",
        "十一、当前局限与后续优化",
        "十二、答辩总结话术",
    ]
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(4)
        run = p.add_run(item)
        set_run_font(run, size=11, color=RGBColor(30, 30, 30))
    doc.add_page_break()


def add_callout(doc: Document, text: str):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_width(table, [6.5])
    cell = table.cell(0, 0)
    shade_cell(cell, LIGHT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run(text)
    set_run_font(run, size=10.5, color=NAVY, bold=True)


def normalize_inline(text: str) -> str:
    text = re.sub(r"`([^`]+)`", r"\1", text)
    text = text.replace("**", "")
    return text


def add_paragraph_with_inline_code(doc: Document, text: str, style=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.space_after = Pt(8)
    parts = re.split(r"(`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = p.add_run(part[1:-1])
            set_run_font(run, name="Consolas", size=10, color=DARK_BLUE)
        else:
            # Very small markdown emphasis support.
            segments = re.split(r"(\*\*[^*]+\*\*)", part)
            for seg in segments:
                if seg.startswith("**") and seg.endswith("**"):
                    run = p.add_run(seg[2:-2])
                    set_run_font(run, bold=True)
                else:
                    run = p.add_run(seg)
                    set_run_font(run)
    return p


def is_table_start(lines, idx):
    return idx + 1 < len(lines) and lines[idx].strip().startswith("|") and re.match(r"^\s*\|?\s*:?-{3,}:?\s*(\|\s*:?-{3,}:?\s*)+\|?\s*$", lines[idx + 1])


def parse_table(lines, idx):
    block = []
    while idx < len(lines) and lines[idx].strip().startswith("|"):
        block.append(lines[idx].strip())
        idx += 1
    header = [c.strip() for c in block[0].strip("|").split("|")]
    rows = [[c.strip() for c in row.strip("|").split("|")] for row in block[2:]]
    return header, rows, idx


def add_markdown_table(doc: Document, header, rows):
    if not header:
        return
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    widths = [6.5 / len(header)] * len(header)
    if len(header) == 2:
        widths = [1.85, 4.65]
    elif len(header) == 3:
        widths = [1.55, 2.45, 2.5]
    set_table_width(table, widths)

    for j, title in enumerate(header):
        cell = table.cell(0, j)
        shade_cell(cell, TABLE_FILL)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(normalize_inline(title))
        set_run_font(run, size=9.5, color=NAVY, bold=True)

    for row in rows:
        cells = table.add_row().cells
        for j, value in enumerate(row[: len(header)]):
            p = cells[j].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(normalize_inline(value))
            set_run_font(run, size=9.2, color=RGBColor(35, 35, 35))
    doc.add_paragraph()


def add_code_block(doc: Document, code_lines):
    for line in code_lines:
        p = doc.add_paragraph(style="CodeBlock")
        p.paragraph_format.left_indent = Inches(0.18)
        p.paragraph_format.right_indent = Inches(0.18)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line if line else " ")
        set_run_font(run, name="Consolas", size=8.5, color=RGBColor(45, 50, 55))
        p_pr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F7F7F7")
        p_pr.append(shd)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(6)


def add_markdown_content(doc: Document, markdown: str):
    lines = markdown.splitlines()
    idx = 0
    in_code = False
    code_lines: list[str] = []
    skip_first_title = True

    while idx < len(lines):
        raw = lines[idx]
        line = raw.rstrip()

        if line.strip().startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            idx += 1
            continue

        if in_code:
            code_lines.append(line)
            idx += 1
            continue

        if not line.strip():
            idx += 1
            continue

        if is_table_start(lines, idx):
            header, rows, idx = parse_table(lines, idx)
            add_markdown_table(doc, header, rows)
            continue

        if line.startswith("# "):
            if skip_first_title:
                skip_first_title = False
            else:
                doc.add_heading(normalize_inline(line[2:].strip()), level=1)
            idx += 1
            continue

        if line.startswith("## "):
            heading = normalize_inline(line[3:].strip())
            doc.add_heading(heading, level=1)
            if heading == "一、项目概述":
                add_callout(doc, "核心定位：以情绪识别为核心，融合机器学习、大模型生成、阅读推荐与可视化手账创作的智能情绪陪伴系统。")
            idx += 1
            continue

        if line.startswith("### "):
            doc.add_heading(normalize_inline(line[4:].strip()), level=2)
            idx += 1
            continue

        if re.match(r"^\d+\.\s+", line.strip()):
            text = re.sub(r"^\d+\.\s+", "", line.strip())
            p = add_paragraph_with_inline_code(doc, text, style="List Number")
            p.paragraph_format.space_after = Pt(4)
            idx += 1
            continue

        if line.strip().startswith("- "):
            text = line.strip()[2:]
            p = add_paragraph_with_inline_code(doc, text, style="List Bullet")
            p.paragraph_format.space_after = Pt(4)
            idx += 1
            continue

        add_paragraph_with_inline_code(doc, line.strip())
        idx += 1


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_document(doc)
    add_cover(doc)
    add_static_toc(doc)
    add_markdown_content(doc, SOURCE.read_text(encoding="utf-8"))
    doc.save(DOCX_OUT)
    print(DOCX_OUT)


if __name__ == "__main__":
    main()
